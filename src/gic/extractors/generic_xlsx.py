"""Generic extractor for Excel files."""

import warnings

import openpyxl

from ..datastructs import AssociatedValues
from .generic import GenericExtractor

# Ignore common openpyxl warning
warnings.filterwarnings(
    "ignore",
    category=UserWarning,
    message="Workbook contains no default style, apply openpyxl's default",
)


class GenericXlsxExtractor(GenericExtractor):
    """Extract data from an Excel file."""

    def extraction_method(self):
        """Extract data from an Excel file."""
        wb = openpyxl.load_workbook(self.input_file)
        if self.sheet_name != "":
            if self.sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet {self.sheet_name} not found in workbook")
            sheet = wb[self.sheet_name]
        else:
            sheet = wb.active

        if sheet is None:
            raise ValueError("No active sheet found in workbook")

        # Extract data after headers
        for row in sheet.iter_rows(min_row=self.start_row):
            data: dict[str, list] = {}
            for column_number, field in self.columns.items():
                index = column_number - 1
                if self.subdelimiter:
                    values = str(row[index].value).split(self.subdelimiter)
                else:
                    values = [str(row[index].value)]
                if field not in data:
                    data[field] = []
                for value in values:
                    value = value.strip()
                    if value not in data[field]:
                        data[field].append(value)
            data = self.cleaner.filter_data(data)
            associated_values = AssociatedValues(data)
            associated_values.import_from_dict(self.constants, overwrite=False)
            associated_values.clean_and_log(self.error_log, row[0].row or 0)
            self.df.append(associated_values)
